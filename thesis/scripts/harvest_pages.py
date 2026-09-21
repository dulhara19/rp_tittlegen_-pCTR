"""Render the NEW thesis to PDF (LibreOffice) and harvest the displayed page
number of every heading / caption / front-matter title into page_map.json.
Re-run whenever content or pagination changes, then re-run format_thesis_docx.py.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

import fitz
from docx import Document

HERE = Path(__file__).resolve().parent
THESIS = HERE.parent
DOCX = THESIS / "Pathiraja_29071_Final_Thesis_NEW.docx"
PDF_DIR = HERE / "_render"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"


def render():
    PDF_DIR.mkdir(exist_ok=True)
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(PDF_DIR),
                    str(DOCX)], check=True, capture_output=True, timeout=600)
    return PDF_DIR / (DOCX.stem + ".pdf")


def norm(s):
    return re.sub(r"[^a-z0-9.:]", "", s.lower())


def page_labels(pdf):
    labels, texts = [], []
    for page in pdf:
        h = page.rect.height
        label = None
        for w in page.get_text("words"):
            if w[1] > h - 70 and re.fullmatch(r"[ivxlc]+|\d+", w[4]):
                label = w[4]
        labels.append(label)
        texts.append(page.get_text())
    return labels, texts


def main():
    pdf_path = render()
    pdf = fitz.open(str(pdf_path))
    labels, texts = page_labels(pdf)
    body_pages = [i for i, l in enumerate(labels) if l and l.isdigit()]
    roman_pages = [i for i, l in enumerate(labels) if l and not l.isdigit()]
    start = body_pages[0]
    doc = Document(str(DOCX))

    out = {"front": {}, "head": {}, "cap": {}}
    # front matter (case-sensitive upper-case titles on roman pages)
    for name in ("Declaration", "Acknowledgement", "Abstract", "Table of Contents",
                 "List of Figures", "List of Tables", "List of Abbreviations"):
        key = name.upper()
        for i in roman_pages:
            if any(line.strip() == key for line in texts[i].splitlines()):
                out["front"][name] = labels[i]
                break
    # headings, monotone through body pages
    ptr = start
    ntexts = {i: norm(texts[i]) for i in body_pages}
    missing = []
    heads = [(p.text.strip()) for p in doc.paragraphs
             if p.style.name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4")]
    for h in heads:
        m = re.match(r"^(\d+(?:\.\d+)*)\s+(.*)$", h)
        num, title = (m.group(1), m.group(2)) if m else ("References", h)
        needle = norm((num + title) if m else title)
        found = None
        for i in body_pages:
            if i < ptr:
                continue
            if needle in ntexts[i]:
                found = i
                break
        if found is None:
            missing.append(h)
            continue
        ptr = found
        out["head"][num] = labels[found]
    # captions
    for p in doc.paragraphs:
        m = re.match(r"^((?:Table|Figure) \d+\.\d+):", p.text)
        if not m:
            continue
        label = m.group(1)
        if label in out["cap"]:
            continue
        needle = norm(label + ":")
        for i in body_pages:
            if needle in ntexts[i]:
                out["cap"][label] = labels[i]
                break
        else:
            missing.append(label)
    (HERE / "page_map.json").write_text(json.dumps(out, indent=1), encoding="utf-8")
    print("pdf pages:", len(pdf), "| body starts at pdf page", start + 1,
          "| front:", out["front"])
    print("headings mapped:", len(out["head"]), "captions mapped:", len(out["cap"]),
          "missing:", missing)


if __name__ == "__main__":
    main()
